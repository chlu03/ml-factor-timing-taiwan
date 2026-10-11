"""Convert the TEJ Excel exports in data/ to CSV files (run once).

Reads only the raw sheets: closing prices, market capitalization, price-to-book ratio,
and the TAIEX (TWII) monthly returns. Everything else is recomputed in Python.

    python -m src.prepare_data
"""
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

DATA = Path(__file__).resolve().parents[1] / 'data'

SHEETS = {                       # output name: (Excel file, sheet)
    'close': ('price.xlsx', '收盤價'),
    'market_cap': ('size.xlsx', '市值'),
    'pbr': ('bm.xlsx', '股價淨值比'),
}


def to_float(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return np.nan


def read_panel(file, sheet):
    """Rows are stocks, columns are months (YYYY/MM). Returns a DataFrame indexed by stock code."""
    wb = openpyxl.load_workbook(DATA / file, read_only=True, data_only=True)
    rows = wb[sheet].iter_rows(values_only=True)
    header = next(rows)
    months = [str(m) for m in header[2:] if m is not None]
    records, codes = [], []
    for row in rows:
        if row[0] is None:
            continue
        codes.append(str(row[0]))
        records.append([to_float(v) for v in row[2:2 + len(months)]])
    wb.close()
    return pd.DataFrame(records, index=pd.Index(codes, name='code'), columns=months)


def read_twii(file='Real.xlsx'):
    """TAIEX monthly returns, labelled by the month in which the return is realized."""
    wb = openpyxl.load_workbook(DATA / file, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = next(r for r in rows if r[0] == '等分')
    twii = next(r for r in rows if r[1] == 'TWII')
    pairs = [(str(m), to_float(v)) for m, v in zip(header[2:], twii[2:]) if m is not None]
    return pd.Series(dict(pairs), name='twii_return').rename_axis('month')


def main():
    for name, (file, sheet) in SHEETS.items():
        panel = read_panel(file, sheet)
        panel.to_csv(DATA / f'{name}.csv')
        print(f'{name}: {panel.shape[0]} stocks x {panel.shape[1]} months -> data/{name}.csv')
    twii = read_twii()
    twii.to_csv(DATA / 'twii.csv')
    print(f'twii: {len(twii)} months -> data/twii.csv')


if __name__ == '__main__':
    main()
